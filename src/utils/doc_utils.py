"""文档工具函数"""
from pathlib import Path
from typing import List, Optional
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from src.utils.logger import get_logger
from src.utils.exceptions import FileProcessingError

logger = get_logger(__name__)


def create_docx(output_path: str, content: str, metadata: dict = None) -> str:
    """创建 DOCX 文件
    
    Args:
        output_path: 输出文件路径
        content: 文档内容
        metadata: 文档元数据
        
    Returns:
        输出文件路径
    """
    try:
        # 创建文档
        doc = Document()
        
        # 设置元数据
        if metadata:
            core_properties = doc.core_properties
            if 'title' in metadata:
                core_properties.title = metadata['title']
            if 'author' in metadata:
                core_properties.author = metadata['author']
            if 'subject' in metadata:
                core_properties.subject = metadata['subject']
        
        # 添加内容
        paragraphs = content.split('\n')
        for paragraph_text in paragraphs:
            if paragraph_text.strip():
                paragraph = doc.add_paragraph(paragraph_text)
                # 设置默认格式
                paragraph.paragraph_format.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
        
        # 确保输出目录存在
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        
        # 保存文档
        doc.save(output_path)
        
        logger.info("docx_created", output=output_path, paragraphs=len(paragraphs))
        return output_path
        
    except Exception as e:
        logger.error("docx_creation_error", output=output_path, error=str(e))
        raise FileProcessingError(f"Failed to create DOCX file: {e}")


def extract_docx_styles(docx_path: str) -> dict:
    """提取 DOCX 文档样式
    
    Args:
        docx_path: DOCX 文件路径
        
    Returns:
        样式信息字典
    """
    try:
        doc = Document(docx_path)
        styles = {}
        
        # 提取段落样式
        paragraph_styles = {}
        for style in doc.styles:
            if style.type == 1:  # 段落样式
                paragraph_styles[style.name] = {
                    'font_size': style.font.size.pt if style.font.size else None,
                    'font_name': style.font.name,
                    'bold': style.font.bold,
                    'italic': style.font.italic
                }
        
        styles['paragraph_styles'] = paragraph_styles
        
        return styles
        
    except Exception as e:
        logger.error("docx_style_extraction_error", file=docx_path, error=str(e))
        raise FileProcessingError(f"Failed to extract DOCX styles: {e}")