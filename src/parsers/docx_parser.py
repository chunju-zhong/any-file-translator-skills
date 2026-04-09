"""DOCX 文件解析器"""
from pathlib import Path
from typing import List, Dict, Any
from docx import Document
from src.parsers.base import BaseParser, ParseResult, Chapter
from src.utils.logger import get_logger
from src.utils.exceptions import FileProcessingError


class DocxParser(BaseParser):
    """DOCX 文件解析器"""
    
    supported_extensions = [".docx"]
    
    def __init__(self):
        super().__init__()
        self.logger = get_logger(__name__)
    
    def parse(self, file_path: str) -> ParseResult:
        """解析 DOCX 文件"""
        try:
            path = Path(file_path)
            
            if not path.exists():
                raise FileProcessingError(f"File not found: {file_path}")
            
            if not self.can_parse(file_path):
                raise FileProcessingError(f"Unsupported file format: {file_path}")
            
            # 使用 python-docx 提取文本
            doc = Document(file_path)
            text = ""
            
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            
            # 提取元数据
            metadata = self._extract_metadata(doc)
            
            self.logger.info(
                "docx_file_parsed",
                file=str(path),
                paragraphs=len(doc.paragraphs),
                size=len(text)
            )
            
            return ParseResult(
                text=text,
                page_count=1,  # DOCX 没有固定页数概念
                metadata={
                    "file_size": path.stat().st_size,
                    "file_name": path.name,
                    "paragraph_count": len(doc.paragraphs),
                    **metadata
                },
                format_type="docx"
            )
        
        except Exception as e:
            self.logger.error("docx_parse_error", file=file_path, error=str(e))
            raise FileProcessingError(f"Failed to parse DOCX file: {e}")
    
    def get_page_count(self, file_path: str) -> int:
        """DOCX 文件返回 1 页"""
        return 1
    
    def extract_chapters(self, file_path: str) -> List[Chapter]:
        """从 DOCX 提取章节信息"""
        chapters = []
        
        try:
            doc = Document(file_path)
            chapter_count = 0
            
            for i, paragraph in enumerate(doc.paragraphs):
                # 简单实现：基于标题样式识别章节
                if paragraph.style.name.startswith('Heading'):
                    chapter_count += 1
                    chapters.append(Chapter(
                        title=paragraph.text.strip(),
                        start_page=1,
                        end_page=1,
                        level=int(paragraph.style.name[-1]) if paragraph.style.name[-1].isdigit() else 1
                    ))
        
        except Exception as e:
            self.logger.warning("docx_chapter_extraction_error", file=file_path, error=str(e))
        
        return chapters
    
    def _extract_metadata(self, doc) -> dict:
        """提取 DOCX 元数据"""
        metadata = {}
        
        try:
            # 提取文档属性
            core_properties = doc.core_properties
            if core_properties.title:
                metadata["title"] = core_properties.title
            if core_properties.author:
                metadata["author"] = core_properties.author
            if core_properties.subject:
                metadata["subject"] = core_properties.subject
            if core_properties.keywords:
                metadata["keywords"] = core_properties.keywords
        
        except Exception as e:
            self.logger.warning("docx_metadata_extraction_error", error=str(e))
        
        return metadata